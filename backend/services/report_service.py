from io import BytesIO
from typing import List, Dict, Any
from datetime import datetime, date
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


class ReportService:
    @staticmethod
    def generate_pdf_report(
        student_data: Dict[str, Any],
        attendance_records: List[Dict[str, Any]],
        summary: Dict[str, Any]
    ) -> bytes:
        """Generate PDF attendance report for a student"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1e40af'),
            alignment=TA_CENTER,
            spaceAfter=20
        )
        
        header_style = ParagraphStyle(
            'CustomHeader',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#1e3a8a'),
            spaceAfter=10
        )

        # Title
        story.append(Paragraph("Attendance Report", title_style))
        story.append(Spacer(1, 0.2 * inch))

        # Student Information
        story.append(Paragraph("Student Information", header_style))
        student_info = [
            ['Name:', student_data.get('full_name', 'N/A')],
            ['Roll Number:', student_data.get('roll_number', 'N/A')],
            ['Department:', student_data.get('department', 'N/A')],
            ['Semester:', str(student_data.get('semester', 'N/A'))],
        ]
        
        student_table = Table(student_info, colWidths=[2 * inch, 3 * inch])
        student_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ]))
        story.append(student_table)
        story.append(Spacer(1, 0.3 * inch))

        # Attendance Summary
        story.append(Paragraph("Attendance Summary", header_style))
        summary_data = [
            ['Total Classes:', str(summary.get('total_classes', 0))],
            ['Present:', str(summary.get('present', 0))],
            ['Absent:', str(summary.get('absent', 0))],
            ['Late:', str(summary.get('late', 0))],
            ['Excused:', str(summary.get('excused', 0))],
            ['Attendance %:', f"{summary.get('percentage', 0)}%"],
        ]
        
        summary_table = Table(summary_data, colWidths=[2 * inch, 3 * inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 0.3 * inch))

        # Attendance Records Table
        if attendance_records:
            story.append(Paragraph("Attendance Records", header_style))
            
            headers = ['Date', 'Subject', 'Status', 'Method']
            data = [headers]
            
            for record in attendance_records:
                marked_at = record.get('marked_at')
                date_str = marked_at.strftime('%Y-%m-%d %H:%M') if marked_at else 'N/A'
                data.append([
                    date_str,
                    record.get('subject_name', 'N/A'),
                    record.get('status', 'N/A'),
                    record.get('method', 'N/A')
                ])
            
            records_table = Table(data, colWidths=[1.5 * inch, 2 * inch, 1 * inch, 1 * inch])
            records_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')]),
            ]))
            story.append(records_table)

        # Footer
        story.append(Spacer(1, 0.5 * inch))
        footer = Paragraph(
            f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles['Normal']
        )
        story.append(footer)

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_excel_report(
        subject_name: str,
        attendance_data: List[Dict[str, Any]]
    ) -> bytes:
        """Generate Excel attendance report for a subject"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        # Define styles
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1E40AF', end_color='1E40AF', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'))
        
        # Write headers
        headers = ['Roll Number', 'Student Name', 'Date', 'Status', 'Method', 'Marked At']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Write data
        for row, record in enumerate(attendance_data, 2):
            marked_at = record.get('marked_at')
            marked_at_str = marked_at.strftime('%Y-%m-%d %H:%M:%S') if marked_at else 'N/A'
            
            ws.cell(row=row, column=1, value=record.get('roll_number', 'N/A'))
            ws.cell(row=row, column=2, value=record.get('student_name', 'N/A'))
            ws.cell(row=row, column=3, value=marked_at_str)
            ws.cell(row=row, column=4, value=record.get('status', 'N/A'))
            ws.cell(row=row, column=5, value=record.get('method', 'N/A'))
            ws.cell(row=row, column=6, value=marked_at_str)
            
            # Apply borders
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = thin_border

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
